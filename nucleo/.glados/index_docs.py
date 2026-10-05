#!/usr/bin/env python3
"""
index_docs.py — Indexa las colecciones del vault para las dos capas de búsqueda.

Las colecciones y el nombre de su carpeta son configurables, no fijos: un vault recién
generado puede no tener todavía ninguna de las carpetas que `config.colecciones()`
declara, y eso no es un error — significa que no hay nada que indexar aún. La
resolución de carpeta por colección vive en `_carpeta_coleccion()`: si el usuario le
dio nombre explícito en `settings.json` → `carpetas`, se usa ese; si no, se acepta
cualquier carpeta de primer nivel cuyo nombre coincida con la colección ignorando un
prefijo numérico (`05_FUENTES` sirve para la colección `fuentes` sin configurar nada),
que es la convención de numeración que ya trae la plantilla de estructura del vault.

Dentro de cada colección, la primera carpeta es el `cliente` (o el proyecto, en
memoria). Se ignora todo lo que empiece por `.`, `~$` o `_`: material no curado o
generado (volcados sin destilar, notas de repos sincronizadas, adjuntos). Formatos:
.pdf, .docx, .md, .txt, .csv. Incremental por mtime; borra de los índices lo que ya no
existe en disco.

Escribe dos índices en la misma pasada, los dos FUERA del vault (`config.data_dir()`):
una base de datos dentro de una carpeta que además sincroniza a la nube se corrompe si
la sincronización la toca a mitad de escritura, y los dos índices pueden pesar más de
lo que debe viajar junto a las notas.
  · vectorstore (Chroma)   → búsqueda semántica, la usa query_docs.py a demanda.
  · lexical.sqlite3 (FTS5) → búsqueda literal, la usa el hook de contexto (milisegundos).

Uso:
    python index_docs.py --vault <ruta>                     # incremental, todo
    python index_docs.py --vault <ruta> --rebuild            # desde cero
    python index_docs.py --vault <ruta> --solo-lexico        # sin cargar el modelo
    python index_docs.py --vault <ruta> --coleccion fuentes  # solo una colección
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config
import lexical

config.force_utf8()

# Rutas que nunca se indexan: temporales de Office, ocultos, y todo lo que empieza por
# "_" (material no curado o generado).
SKIP_PREFIXES = ("~$", ".", "_")
EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".csv"}
CHUNK_WORDS = 400
OVERLAP_WORDS = 50
NOMBRE_COLECCION_VECTORES = "docs"

# Modelo multilingüe por defecto si `settings.json` no fija `modelo_embeddings`: los
# documentos funcionales de un vault en español necesitan un modelo que entienda
# español, no uno solo de inglés, que devolvería vecindades sin sentido en silencio.
MODELO_EMBEDDINGS_DEFECTO = "paraphrase-multilingual-MiniLM-L12-v2"


def chunk_text(text: str) -> list[str]:
    words = text.split()
    step = CHUNK_WORDS - OVERLAP_WORDS
    return [" ".join(words[i:i + CHUNK_WORDS]) for i in range(0, len(words), step) if words[i:i + CHUNK_WORDS]]


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if ext == ".docx":
        import docx
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    return path.read_text(encoding="utf-8", errors="ignore")


def client_of(rel: Path) -> str:
    return rel.parts[0] if len(rel.parts) > 1 else "_general"


def _carpeta_coleccion(raiz: Path, nombre: str) -> Path | None:
    """Carpeta del vault para esta colección, o None si no hay ninguna — no es un error.

    Prioridad: el nombre que el usuario le dio en `settings.json` → `carpetas` (el
    mismo diccionario que usa `context_injector.py` para roles como `clientes` o
    `proyectos`). Sin configurar, se acepta una carpeta de primer nivel cuyo nombre,
    quitando un posible prefijo numérico (`05_`, `07_`...), coincida con la colección.
    """
    configurada = config.carpetas().get(nombre)
    if configurada:
        return raiz / configurada
    if not raiz.exists():
        return None
    objetivo = nombre.lower()
    for hijo in sorted(raiz.iterdir()):
        if not hijo.is_dir():
            continue
        limpio = re.sub(r"^\d+_", "", hijo.name).lower()
        if limpio == objetivo:
            return hijo
    return None


def candidate_files(root: Path) -> list[Path]:
    """Ficheros indexables bajo `root`, saltando rutas con prefijo ignorado."""
    out = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in EXTENSIONS:
            continue
        if any(part.startswith(SKIP_PREFIXES) for part in path.relative_to(root).parts):
            continue
        out.append(path)
    return out


def indexed_state_vector(collection, coleccion: str) -> dict[str, float]:
    """{source: mtime} de lo que ya está en el índice semántico para esa colección."""
    res = collection.get(where={"coleccion": coleccion}, include=["metadatas"], limit=1_000_000)
    return {m["source"]: round(m.get("mtime", 0), 3) for m in res["metadatas"] or []}


def indexar(raiz: Path, solo_lexico: bool = False, coleccion: str | None = None,
            rebuild: bool = False) -> int:
    """Indexa `raiz` (la raíz del vault) en las colecciones configuradas.

    Devuelve el número de documentos indexados (nuevos o cambiados) en esta pasada.
    """
    nombres = [n for n in config.colecciones() if not coleccion or n == coleccion]
    if not nombres:
        print(f"Colección desconocida: {coleccion}. Hay: " + ", ".join(config.colecciones()))
        return 0

    collection = None
    if not solo_lexico:
        import chromadb
        from chromadb.utils import embedding_functions
        vectorstore_path = config.data_dir() / "vectorstore"
        vectorstore_path.mkdir(parents=True, exist_ok=True)
        client = chromadb.PersistentClient(path=str(vectorstore_path))
        if rebuild:
            try:
                client.delete_collection(NOMBRE_COLECCION_VECTORES)
            except Exception:
                pass
        modelo = config.local_settings().get("modelo_embeddings") or MODELO_EMBEDDINGS_DEFECTO
        ef = embedding_functions.SentenceTransformerEmbeddingFunction(model_name=modelo)
        collection = client.get_or_create_collection(NOMBRE_COLECCION_VECTORES, embedding_function=ef)

    if rebuild and lexical.lexical_db_path().exists():
        lexical.lexical_db_path().unlink()

    total_i = total_s = total_r = 0
    for nombre in nombres:
        root = _carpeta_coleccion(raiz, nombre)
        if not root or not root.exists():
            print(f"[{nombre}] sin carpeta configurada o inexistente todavía. Salto.")
            continue
        print(f"[{nombre}] {root}")
        files = candidate_files(root)
        state_vec = indexed_state_vector(collection, nombre) if collection is not None else {}
        state_lex = lexical.indexed_state(nombre)
        on_disk = set()
        indexed = skipped = removed = 0

        for path in files:
            rel = path.relative_to(root)
            source = rel.as_posix()
            on_disk.add(source)
            mtime = round(path.stat().st_mtime, 3)
            if (collection is None or state_vec.get(source) == mtime) and state_lex.get(source) == mtime:
                skipped += 1
                continue
            try:
                chunks = chunk_text(extract_text(path))
            except Exception as e:
                print(f"  [ERROR] {source}: {e}")
                continue
            if not chunks:
                print(f"  [WARN] {source}: sin texto extraíble (¿PDF escaneado?)")
                continue
            cliente = client_of(rel)
            print(f"  [INDEX] {source} ({len(chunks)} fragmentos, cliente={cliente})")
            lexical.replace_file(source, nombre, cliente, mtime, chunks)
            if collection is not None:
                if source in state_vec:
                    collection.delete(where={"$and": [{"source": source}, {"coleccion": nombre}]})
                collection.add(
                    documents=chunks,
                    ids=[f"{nombre}::{source}::{i}" for i in range(len(chunks))],
                    metadatas=[{"source": source, "coleccion": nombre, "cliente": cliente,
                                "chunk": i, "mtime": mtime} for i in range(len(chunks))],
                )
            indexed += 1

        for source in (set(state_vec) | set(state_lex)) - on_disk:
            lexical.forget_file(source, nombre)
            if collection is not None and source in state_vec:
                collection.delete(where={"$and": [{"source": source}, {"coleccion": nombre}]})
            removed += 1

        print(f"  → indexados {indexed} | sin cambios {skipped} | eliminados {removed}")
        total_i, total_s, total_r = total_i + indexed, total_s + skipped, total_r + removed

    capas = "léxica" if solo_lexico else "léxica + semántica"
    print(f"\nCapas: {capas} | Total indexados: {total_i} | Sin cambios: {total_s} | Eliminados: {total_r}")
    filas = lexical.stats()
    if filas:
        print("\nÍndice léxico:")
        for col, nf, nc in filas:
            print(f"  {col:<10} {nf:>4} ficheros  {nc:>6} fragmentos")
    return total_i


def main() -> int:
    ap = argparse.ArgumentParser(description="Indexa las colecciones del vault para las "
                                              "dos capas de búsqueda.")
    ap.add_argument("--vault", required=True,
                     help="Raíz del vault a indexar. Obligatorio: sin esto solo se "
                          "podría probar esta herramienta contra el vault real.")
    ap.add_argument("--rebuild", action="store_true")
    ap.add_argument("--solo-lexico", action="store_true",
                     help="No carga el modelo de embeddings: solo el índice literal.")
    ap.add_argument("--coleccion", help="Indexar solo esta colección.")
    args = ap.parse_args()

    indexar(Path(args.vault), solo_lexico=args.solo_lexico, coleccion=args.coleccion,
            rebuild=args.rebuild)
    return 0


if __name__ == "__main__":
    sys.exit(main())
